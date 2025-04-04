from flask import Flask, render_template, request, jsonify, session
import os
import docx  # python-docx for Word files
import google.generativeai as genai
import pandas as pd
import pypdf as pdf
from pypdf import PdfReader
import json
import re
import uuid
from difflib import get_close_matches
from datetime import datetime
import requests
from bs4 import BeautifulSoup

app = Flask(__name__)
app.secret_key = "your_secret_key_here"

# Configure Gemini API
GENAI_API_KEY = "AIzaSyC7W2hDNU9k5S52d5vlGk2uwRKFSIXQnO4"  # Replace with your actual API key
genai.configure(api_key=GENAI_API_KEY)
model = genai.GenerativeModel("gemini-2.0-flash")

# Load chatbot responses from JSON
with open(r"C:\Users\Chethan SM\OneDrive\Desktop\ahana chatbot\config.json", "r", encoding="utf-8") as f:
    data = json.load(f)

company_website = data.get("company_websites", {})
responses = data.get("responses", {})

# Define paths
EXCEL_FILE = "chat-history.xlsx"
UPLOAD_FOLDER = "uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
ALLOWED_EXTENSIONS = {"pdf", "txt", "docx", "csv", "xlsx"}

# Store extracted file data
file_text_data = {}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_file(file_path, file_ext):
    text = ""
    if file_ext == "pdf":
        doc = PdfReader(file_path)
        for page in doc.pages:
            text += page.extract_text() + "\n"
    elif file_ext == "docx":
        doc = docx.Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs])
    elif file_ext == "txt":
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
    elif file_ext in ["csv", "xlsx"]:
        df = pd.read_excel(file_path) if file_ext == "xlsx" else pd.read_csv(file_path)
        text = df.to_string(index=False)

    file_text_data[file_path] = text
    return text

def get_session_id():
    if "session_id" not in session:
        session["session_id"] = str(uuid.uuid4())
    return session["session_id"]

def save_chat_to_excel(user_message, bot_response):
    session_id = get_session_id()
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    chat_entry = pd.DataFrame([[session_id, timestamp, user_message, bot_response]],
                              columns=["Session ID", "Timestamp", "User Message", "Bot Response"])
    if not os.path.exists(EXCEL_FILE):
        chat_entry.to_excel(EXCEL_FILE, index=False)
    else:
        with pd.ExcelWriter(EXCEL_FILE, mode='a', engine='openpyxl', if_sheet_exists='overlay') as writer:
            chat_entry.to_excel(writer, index=False, header=False, startrow=writer.sheets["Sheet1"].max_row)

@app.route('/')
def home():
    session.pop("session_id", None)
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file selected"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400
    if file and allowed_file(file.filename):
        file_ext = file.filename.rsplit('.', 1)[1].lower()
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
        file.save(file_path)
        extracted_text = extract_text_from_file(file_path, file_ext)
        file_text_data[file.filename] = extracted_text

        response_text = f"File '{file.filename}' uploaded successfully! Here is the extracted text:\n\n{extracted_text[:500]}..."
        save_chat_to_excel(f"Uploaded file: {file.filename}", response_text)

        return jsonify({"response": response_text, "file_name": file.filename, "extracted_text": extracted_text})
    
    return jsonify({"error": "File type not allowed"}), 400

# Web Scraping API
@app.route('/scrape', methods=['POST'])
def web_scraper():
    data = request.get_json()
    url = data.get("url", "").strip()
    
    if not url:
        return jsonify({"error": "No URL provided"}), 400
    
    scraped_data = scrape_website(url)
    return jsonify({"response": scraped_data})

def scrape_website(url):
    headers = {"User-Agent": "Mozilla/5.0"}
    try:
        response = requests.get(url, headers=headers, timeout=10)
        if response.status_code != 200:
            return "Failed to retrieve data. The website might be blocking scraping."

        soup = BeautifulSoup(response.text, "html.parser")

        headlines = [h.text.strip() for h in soup.find_all("h2")[:5]]
        content = " ".join([p.text.strip() for p in soup.find_all("p")[:5]])

        if headlines:
            return "Top headlines:\n" + "\n".join(headlines)
        elif content:
            return "Main content:\n" + content
        else:
            return "No relevant data found on the page."

    except requests.exceptions.RequestException as e:
        return f"Error fetching data: {e}"

@app.route('/chatbot', methods=['POST'])
def chatbot():
    data = request.get_json()
    user_message = data.get("message", "").strip().lower()
    
    if not user_message:
        response_text = "Hi! I am Ahana, your chatbot. How can I assist you today?"
        save_chat_to_excel(user_message, response_text)
        return jsonify({"response": response_text})

    for file_name, file_text in file_text_data.items():
        if file_name.lower() in user_message:
            response_text = f"Here is the extracted text from {file_name}:\n\n{file_text[:500]}..."
            save_chat_to_excel(user_message, response_text)
            return jsonify({"response": response_text})

    for keyword, url in company_website.items():
        if re.search(rf"\b{keyword}\b", user_message, re.IGNORECASE):
            response_text = f"Redirecting to {keyword} website..."
            if keyword == "support":
                response_text += " Please contact support@ahanait.com for more help."
            save_chat_to_excel(user_message, response_text)
            return jsonify({"response": response_text, "url": url})
    
    for keyword, response_text in responses.items():
        if re.search(rf"\b{keyword}\b", user_message, re.IGNORECASE):
            save_chat_to_excel(user_message, response_text)
            return jsonify({"response": response_text})

    if "scrape" in user_message or "fetch from" in user_message:
        match = re.search(r"https?://[^\s]+", user_message)
        if match:
            url = match.group()
            scraped_data = scrape_website(url)
            return jsonify({"response": f"Here is the data from {url}:\n\n{scraped_data}"})

    try:
        gemini_response = model.generate_content(user_message)
        response_text = gemini_response.text if gemini_response else "I'm not sure, can you rephrase?"
    except Exception as e:
        response_text = "I'm having trouble understanding. Please try again later."

    save_chat_to_excel(user_message, response_text)
    return jsonify({"response": response_text})

if __name__ == '__main__':
    app.run(debug=True)